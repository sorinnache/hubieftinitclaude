"""Istoric stocuri: în ultima zi a fiecărei luni (ora 23:55, ora României) se salvează automat
situația stocului de marfă (produs, cantitate, preț de achiziție fără TVA, valoare) și se poate
descărca oricând ca fișier Excel pentru contabil."""
import asyncio
import calendar
import io
import logging
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse

from . import audit
from .auth import get_current_user, require_admin
from .db import get_db, new_id, now_iso
from .stock import NOT_BUNDLE, STOCK_STATUSES, get_settings

log = logging.getLogger("stock_reports")
router = APIRouter(prefix="/stock/snapshots", tags=["Istoric stocuri"])
TZ = ZoneInfo("Europe/Bucharest")
COMPANY = "IEFTINIT SRL"
RUN_AT = (23, 55)  # ora la care se face situația în ultima zi a lunii

LUNI = ["ianuarie", "februarie", "martie", "aprilie", "mai", "iunie", "iulie", "august", "septembrie",
        "octombrie", "noiembrie", "decembrie"]


def _period_label(period: str) -> str:
    y, m = period.split("-")
    return f"{LUNI[int(m) - 1]} {y}"


def _last_day(y: int, m: int) -> int:
    return calendar.monthrange(y, m)[1]


async def create_snapshot(period: str, kind: str = "auto", actor: str = "Sistem", note: str = "") -> dict:
    db = get_db()
    vs = await db.stock_variants.find({"status": STOCK_STATUSES, "deleted": {"$ne": True}, "untracked": {"$ne": True},
                                       "on_hand": {"$gt": 0}, **NOT_BUNDLE}, {"_id": 0}) \
        .sort([("product_title", 1), ("variant_title", 1)]).to_list(20000)
    rows = []
    for v in vs:
        cost = v.get("cost")
        cost = round(float(cost), 2) if cost is not None and float(cost) > 0 else None
        qty = int(v.get("on_hand") or 0)
        rows.append({"variant_id": v["variant_id"], "product_title": v.get("product_title") or "",
                     "variant_title": v.get("variant_title") or "", "sku": v.get("sku") or "",
                     "qty": qty, "cost": cost, "value": round(qty * cost, 2) if cost is not None else None})
    now = datetime.now(TZ)
    snap = {
        "id": new_id(), "period": period, "kind": kind, "note": note,
        "as_of": now.isoformat(), "created_at": now_iso(), "created_by": actor,
        "rows": rows, "products": len(rows), "units": sum(r["qty"] for r in rows),
        "value": round(sum(r["value"] or 0 for r in rows), 2),
        "missing_cost": sum(1 for r in rows if r["cost"] is None),
    }
    await db.stock_snapshots.insert_one(dict(snap))
    await audit.log(resource_type="stock", resource_id=snap["id"], action="stock.snapshot", actor=None,
                    actor_name=actor, details={"period": period, "kind": kind, "products": snap["products"],
                                               "value": snap["value"]})
    return snap


async def monthly_check():
    """Rulează periodic: în ultima zi a lunii după 23:55 face situația lunii (o singură dată).
    Dacă serverul a fost oprit atunci, o face imediat ce pornește, marcată „generată cu întârziere”."""
    db = get_db()
    s = await get_settings()
    if not s.get("start_at"):
        return  # stocul nu e încă urmărit (fără inventar inițial)
    now = datetime.now(TZ)
    period = now.strftime("%Y-%m")
    if now.day == _last_day(now.year, now.month) and (now.hour, now.minute) >= RUN_AT:
        if not await db.stock_snapshots.find_one({"period": period, "kind": {"$in": ["auto", "late"]}}):
            await create_snapshot(period, "auto")
        return
    # situația lunii trecute lipsește? (server oprit în ultima zi) — o facem acum, marcată
    prev = (now.replace(day=1) - timedelta(days=1))
    pperiod = prev.strftime("%Y-%m")
    started = datetime.fromisoformat(s["start_at"].replace("Z", "+00:00")).astimezone(TZ)
    if started.strftime("%Y-%m") <= pperiod and \
            not await db.stock_snapshots.find_one({"period": pperiod, "kind": {"$in": ["auto", "late"]}}):
        await create_snapshot(pperiod, "late", note=f"Generată cu întârziere, pe {now.strftime('%d.%m.%Y %H:%M')} "
                                                    f"(serverul nu a rulat la final de lună)")


async def snapshot_loop():
    await asyncio.sleep(60)
    while True:
        try:
            await monthly_check()
        except Exception as e:  # noqa: BLE001
            log.error("situație lunară stoc: %s", e)
        await asyncio.sleep(120)


# ------------------------------------------------------------------ API
@router.get("")
async def list_snapshots(_: dict = Depends(get_current_user)):
    rows = await get_db().stock_snapshots.find({}, {"_id": 0, "rows": 0}).sort([("period", -1), ("created_at", -1)]) \
        .to_list(500)
    for r in rows:
        r["label"] = _period_label(r["period"])
    s = await get_settings()
    now = datetime.now(TZ)
    nxt = now.replace(day=_last_day(now.year, now.month), hour=RUN_AT[0], minute=RUN_AT[1], second=0, microsecond=0)
    if now >= nxt:
        y, m = (now.year + (now.month == 12), now.month % 12 + 1)
        nxt = nxt.replace(year=y, month=m, day=_last_day(y, m))
    return {"items": rows, "next_run": nxt.isoformat(), "started": bool(s.get("start_at"))}


@router.post("")
async def manual_snapshot(request: Request, user: dict = Depends(require_admin)):
    now = datetime.now(TZ)
    snap = await create_snapshot(now.strftime("%Y-%m"), "manual", actor=user["name"],
                                 note=f"Generată manual pe {now.strftime('%d.%m.%Y %H:%M')}")
    snap.pop("rows", None)
    return snap


@router.get("/{sid}/xlsx")
async def download_xlsx(sid: str, _: dict = Depends(get_current_user)):
    snap = await get_db().stock_snapshots.find_one({"id": sid}, {"_id": 0})
    if not snap:
        raise HTTPException(404, "Situație inexistentă")
    data = build_xlsx(snap)
    fname = f"Stoc_marfa_IEFTINIT_{snap['period']}{'_manual' if snap['kind'] == 'manual' else ''}.xlsx"
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


def build_xlsx(snap: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    as_of = datetime.fromisoformat(snap["as_of"]).astimezone(TZ)
    wb = Workbook()
    ws = wb.active
    ws.title = "Stoc marfă"
    bold = Font(name="Arial", bold=True, size=10)
    normal = Font(name="Arial", size=10)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", start_color="D9E1F2")
    warn_fill = PatternFill("solid", start_color="FCE4D6")
    money = '#,##0.00'

    ws["A1"] = COMPANY
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws["A2"] = f"Situația stocului de marfă la {as_of.strftime('%d.%m.%Y')} (ora {as_of.strftime('%H:%M')})"
    ws["A2"].font = Font(name="Arial", bold=True, size=11)
    ws["A3"] = "Valori la prețul de achiziție, fără TVA, în lei. Doar produsele cu stoc în depozit."
    ws["A3"].font = Font(name="Arial", italic=True, size=9, color="595959")
    r0 = 4
    if snap.get("note"):
        ws.cell(row=r0, column=1, value=snap["note"]).font = Font(name="Arial", italic=True, size=9, color="C00000")
        r0 += 1
    if snap.get("missing_cost"):
        ws.cell(row=r0, column=1, value=(f"Atenție: {snap['missing_cost']} produse nu au" if snap['missing_cost'] > 1 else "Atenție: 1 produs nu are")
                                        + " preț de achiziție setat (rânduri marcate) și nu e inclus în valoarea totală.").font = \
            Font(name="Arial", bold=True, size=9, color="C00000")
        r0 += 1

    hr = r0 + 1
    headers = ["Nr. crt.", "Denumire produs", "Variantă", "Cod (SKU)", "U.M.", "Cantitate",
               "Preț achiziție fără TVA (lei/buc)", "Valoare fără TVA (lei)"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font, cell.fill, cell.border = bold, head_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hr].height = 32

    first = hr + 1
    for i, r in enumerate(snap["rows"]):
        row = first + i
        vals = [i + 1, r["product_title"], r["variant_title"], r["sku"], "buc", r["qty"], r["cost"], None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font, cell.border = normal, border
        ws.cell(row=row, column=8, value=f"=F{row}*G{row}" if r["cost"] is not None else None)
        ws.cell(row=row, column=8).font, ws.cell(row=row, column=8).border = normal, border
        for c in (7, 8):
            ws.cell(row=row, column=c).number_format = money
        ws.cell(row=row, column=6).number_format = '#,##0'
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if r["cost"] is None:
            for c in range(1, 9):
                ws.cell(row=row, column=c).fill = warn_fill
            ws.cell(row=row, column=7, value="lipsă preț").font = Font(name="Arial", size=10, color="C00000")
    last = first + len(snap["rows"]) - 1
    tr = last + 1
    ws.cell(row=tr, column=2, value="TOTAL").font = bold
    ws.cell(row=tr, column=6, value=f"=SUM(F{first}:F{last})" if snap["rows"] else 0).font = bold
    ws.cell(row=tr, column=8, value=f"=SUM(H{first}:H{last})" if snap["rows"] else 0).font = bold
    ws.cell(row=tr, column=6).number_format = '#,##0'
    ws.cell(row=tr, column=8).number_format = money
    for c in range(1, 9):
        ws.cell(row=tr, column=c).border = border
        ws.cell(row=tr, column=c).fill = head_fill

    widths = [8, 60, 20, 22, 6, 11, 18, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = ws.cell(row=first, column=1)
    if snap["rows"]:
        ws.auto_filter.ref = f"A{hr}:H{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    ws.oddFooter.right.text = "Pagina &P din &N"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
